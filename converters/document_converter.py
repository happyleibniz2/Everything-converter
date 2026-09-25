"""PDF and Office/Archive converters (Bronze rank and above).

These wrap optional third-party libraries. Import failures are expected on a
bare install, so each converter reports ``is_available`` instead of exploding;
the registry filters unavailable converters out at import time, which keeps the
Formats page honest about what this machine can actually do.
"""

import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Tuple

from converters.base import Converter

CREATE_NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0

IMAGE_EXTENSIONS_DEFAULT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp")

OFFICE_INPUT_EXTENSIONS = (".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp")
ARCHIVE_INPUT_EXTENSIONS = (".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".xz")


def _module(name: str) -> bool:
    try:
        __import__(name)
        return True
    except Exception:
        return False


class DocumentConverter(Converter):
    """Base for document-family converters that may depend on missing libs."""

    category = "Document"

    def is_available(self) -> bool:  # pragma: no cover - overridden
        return True

    def convert(self, input_file, output_file):
        raise NotImplementedError


class PdfToImageConverter(DocumentConverter):
    """PDF → PNG/JPG via PyMuPDF (fitz), one image per first page."""

    def __init__(self, name: str, output_extension: str):
        self.name = name
        self.input_extensions: Tuple[str, ...] = (".pdf",)
        self.output_extension = output_extension

    def is_available(self) -> bool:
        return _module("fitz")

    def convert(self, input_file, output_file):
        import fitz  # PyMuPDF

        with fitz.open(input_file) as document:
            page = document.load_page(0)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            if self.output_extension in (".jpg", ".jpeg"):
                pixmap.pil_save(output_file, format="JPEG", quality=92)
            else:
                pixmap.save(output_file)


class ImageToPdfConverter(DocumentConverter):
    """Any Pillow-readable image → PDF."""

    name = "Images → PDF"

    def __init__(self, input_extensions: Tuple[str, ...] = IMAGE_EXTENSIONS_DEFAULT):
        self.input_extensions = tuple(input_extensions)
        self.output_extension = ".pdf"

    def is_available(self) -> bool:
        return _module("PIL")

    def convert(self, input_file, output_file):
        from PIL import Image

        image = Image.open(input_file)
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        image.save(output_file, format="PDF")


class PdfToTextConverter(DocumentConverter):
    """PDF → TXT via pypdf."""

    name = "PDF → TXT"
    output_extension = ".txt"

    def __init__(self, input_extensions: Tuple[str, ...] = (".pdf",)):
        self.input_extensions = tuple(input_extensions)

    def is_available(self) -> bool:
        return _module("pypdf")

    def convert(self, input_file, output_file):
        from pypdf import PdfReader

        reader = PdfReader(str(input_file))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        Path(output_file).write_text(text.strip(), encoding="utf-8")


class OfficeToPdfConverter(DocumentConverter):
    """Office documents → PDF using LibreOffice when installed."""

    name = "Office → PDF"

    def __init__(self, input_extensions: Tuple[str, ...] = OFFICE_INPUT_EXTENSIONS):
        self.input_extensions = tuple(input_extensions)
        self.output_extension = ".pdf"

    @staticmethod
    def _soffice():
        return shutil.which("soffice") or shutil.which("libreoffice")

    def is_available(self) -> bool:
        return bool(self._soffice())

    def convert(self, input_file, output_file):
        soffice = self._soffice()
        target_dir = str(Path(output_file).parent)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", target_dir, input_file],
            check=True, capture_output=True, timeout=180, creationflags=CREATE_NO_WINDOW,
        )
        produced = Path(target_dir) / (Path(input_file).stem + ".pdf")
        if produced.resolve() != Path(output_file).resolve():
            produced.replace(output_file)


class OfficeToTxtConverter(DocumentConverter):
    """DOCX/XLSX/PPTX → TXT with python-docx / openpyxl / python-pptx."""

    def __init__(self, extension: str, module: str):
        self.input_extensions: Tuple[str, ...] = (extension,)
        self.output_extension = ".txt"
        self.module = module
        self.name = f"{extension[1:].upper()} \u2192 TXT"

    def is_available(self) -> bool:
        return _module(self.module)

    def convert(self, input_file, output_file):
        extension = self.input_extensions[0]
        lines = []
        if extension == ".docx":
            import docx

            document = docx.Document(input_file)
            lines = [paragraph.text for paragraph in document.paragraphs]
        elif extension == ".xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(input_file, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                lines.append(f"# {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    lines.append("\t".join("" if cell is None else str(cell) for cell in row))
        elif extension == ".pptx":
            from pptx import Presentation

            presentation = Presentation(input_file)
            for index, slide in enumerate(presentation.slides, 1):
                lines.append(f"# Slide {index}")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        lines.append(shape.text_frame.text)
        Path(output_file).write_text("\n".join(lines).strip(), encoding="utf-8")


class ArchiveListConverter(DocumentConverter):
    """Archives → a TXT listing of their contents (ZIP/TAR natively, others by lib)."""

    def __init__(self, extension: str):
        self.input_extensions: Tuple[str, ...] = (extension,)
        self.output_extension = ".txt"
        self.name = f"{extension[1:].upper()} \u2192 Listing"

    def is_available(self) -> bool:
        extension = self.input_extensions[0]
        if extension in (".zip", ".tar", ".gz", ".tgz", ".bz2", ".xz"):
            return True
        if extension == ".7z":
            return _module("py7zr")
        if extension == ".rar":
            return _module("rarfile")
        return False

    def convert(self, input_file, output_file):
        extension = self.input_extensions[0]
        names = []
        if extension == ".zip":
            with zipfile.ZipFile(input_file) as archive:
                names = archive.namelist()
        elif extension in (".tar", ".gz", ".tgz", ".bz2", ".xz"):
            import tarfile

            mode = "r:*"
            with tarfile.open(input_file, mode) as archive:
                names = archive.getnames()
        elif extension == ".7z":
            import py7zr

            with py7zr.SevenZipFile(input_file) as archive:
                names = archive.getnames()
        elif extension == ".rar":
            import rarfile

            with rarfile.RarFile(input_file) as archive:
                names = archive.namelist()
        Path(output_file).write_text("\n".join(names), encoding="utf-8")


IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp")

DOCUMENT_CONVERTERS = [
    PdfToImageConverter("PDF → PNG", ".png"),
    PdfToImageConverter("PDF → JPG", ".jpg"),
    PdfToTextConverter(),
    ImageToPdfConverter(IMAGE_EXTENSIONS_DEFAULT),
    OfficeToPdfConverter(OFFICE_INPUT_EXTENSIONS),
    OfficeToTxtConverter(".docx", "docx"),
    OfficeToTxtConverter(".xlsx", "openpyxl"),
    OfficeToTxtConverter(".pptx", "pptx"),
    *[ArchiveListConverter(ext) for ext in ARCHIVE_INPUT_EXTENSIONS],
]

# Only ship converters whose dependencies exist; the Formats page reads this.
AVAILABLE_DOCUMENT_CONVERTERS = [c for c in DOCUMENT_CONVERTERS if c.is_available()]
